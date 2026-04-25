#!/usr/bin/env python3
"""
build_corpus.py — Assemble the full 2010-2025 classification corpus.

Sources
-------
1. tgIITulua/processed-data/datosTG.xlsx
   Expert-labeled gold set: 2010-2018, N=113 rows (111 with category labels).
   Columns used: ID, NOMBRE DEL TRABAJO, PALABRAS CLAVE, RESUMEN,
                 CATEGORÍA, CATEGORIA REVISA.

2. versions/latest/Seguimiento TG1 -TG2 - C1 (Autoguardado).xlsx
   Program tracking file: 2018-2025, N=131 projects.
   Sheets used: 2019-2024 (2018 already covered by datosTG; sheet '2024'
   contains the 2025A cohort).
   Columns: No., Propuesta/Proyecto, Estudiante(s), [Año,] Resumen,
            Palabras clave[s|vs], ...

Output
------
    papers/tg-roger-joshua/sections/_auto/corpus.xlsx
    papers/tg-roger-joshua/sections/_auto/corpus_stats.txt

USAGE
    bash containers/run.sh paper-data \\
        python papers/tg-roger-joshua/scripts/build_corpus.py
"""

from __future__ import annotations
import sys
from pathlib import Path

try:
    import pandas as pd
    import openpyxl  # noqa: F401 – checked at import time
except ImportError as e:
    print(f"ERROR: missing dep ({e}).", file=sys.stderr)
    sys.exit(2)

PAPER = Path(__file__).resolve().parent.parent
TGII  = PAPER / "sources" / "tgIITulua"

DATOS_XLSX     = TGII / "processed-data" / "datosTG.xlsx"
SEGUIMIENTO_XLSX = (
    PAPER / "versions" / "latest"
    / "Seguimiento TG1 -TG2 - C1 (Autoguardado).xlsx"
)
OUT_DIR = PAPER / "sections" / "_auto"
OUT_XLSX = OUT_DIR / "corpus.xlsx"
OUT_STATS = OUT_DIR / "corpus_stats.txt"

# Sheets in Seguimiento that are NOT already covered by datosTG.
# Sheet '2018' projects overlap with datosTG 2018 entries → skip it.
SEGUIMIENTO_YEARS = ["2019", "2020", "2021", "2022", "2023", "2024"]


# -----------------------------------------------------------------------
# Loaders
# -----------------------------------------------------------------------

def load_datos() -> pd.DataFrame:
    """Load the datosTG gold set."""
    df = pd.read_excel(DATOS_XLSX)
    df.columns = [c.strip() for c in df.columns]
    rename = {
        "ID": "id",
        "NOMBRE DEL TRABAJO": "title",
        "PALABRAS CLAVE": "keywords",
        "RESUMEN": "abstract",
        "CATEGORÍA": "auto_category",
        "CATEGORIA REVISA": "expert_category",
    }
    present = {k: v for k, v in rename.items() if k in df.columns}
    df = df[list(present)].rename(columns=present)
    df = df.dropna(subset=["id"]).reset_index(drop=True)
    df["id"] = df["id"].astype(str).str.strip()
    for col in ("title", "keywords", "abstract"):
        if col in df.columns:
            df[col] = df[col].fillna("").astype(str)
    df["source"] = "datosTG"
    return df


def _col_idx(header: list, *candidates: str) -> int | None:
    """Return index of first header matching any candidate (case-insensitive)."""
    hl = [str(h).lower().strip() if h else "" for h in header]
    for cand in candidates:
        cl = cand.lower()
        for i, h in enumerate(hl):
            if cl in h:
                return i
    return None


def load_seguimiento() -> pd.DataFrame:
    """Load 2019-2025 projects from the Seguimiento file."""
    wb = openpyxl.load_workbook(SEGUIMIENTO_XLSX, read_only=True, data_only=True)
    records = []

    for sheet_name in SEGUIMIENTO_YEARS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        # Locate the column-header row (contains 'No.' in first cell)
        hdr_idx = next(
            (i for i, r in enumerate(rows) if str(r[0]).strip() in ("No.", "No")),
            None,
        )
        if hdr_idx is None:
            continue
        header = list(rows[hdr_idx])

        # Map column names → indices
        title_col   = _col_idx(header, "propuesta", "proyecto")
        abstract_col = _col_idx(header, "resumen")
        kw_col       = _col_idx(header, "clave", "clves")
        year_col     = _col_idx(header, "año", "ano")

        # Data rows: first column is an integer (project number)
        data_rows = [r for r in rows[hdr_idx + 1:]
                     if isinstance(r[0], (int, float)) and r[0] is not None
                     and int(r[0]) == r[0]]

        for r in data_rows:
            seq = int(r[0])
            title    = str(r[title_col]).strip()   if title_col    is not None and r[title_col]    else ""
            abstract = str(r[abstract_col]).strip() if abstract_col is not None and r[abstract_col] else ""
            keywords = str(r[kw_col]).strip()       if kw_col       is not None and r[kw_col]       else ""
            # Year: use sheet-label year if column missing
            if year_col is not None and r[year_col] and str(r[year_col]).isdigit():
                year = int(r[year_col])
            else:
                year = int(sheet_name)
            # Sheet '2024' = 2025A cohort → label as 2025
            display_year = 2025 if sheet_name == "2024" else year
            pid = f"{display_year}S{seq:02d}"
            records.append({
                "id":              pid,
                "title":           title,
                "keywords":        keywords,
                "abstract":        abstract,
                "auto_category":   None,
                "expert_category": None,
                "source":          f"seguimiento_{sheet_name}",
            })

    return pd.DataFrame(records)


# -----------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------

def main() -> int:
    if not DATOS_XLSX.exists():
        print(f"ERROR: {DATOS_XLSX} not found", file=sys.stderr)
        return 2
    if not SEGUIMIENTO_XLSX.exists():
        print(f"ERROR: {SEGUIMIENTO_XLSX} not found", file=sys.stderr)
        return 2

    datos = load_datos()
    seg   = load_seguimiento()

    corpus = pd.concat([datos, seg], ignore_index=True)
    corpus = corpus.fillna("")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    corpus.to_excel(OUT_XLSX, index=False)

    # Stats
    lines = [
        f"Corpus assembled: {len(corpus)} projects",
        f"  datosTG (gold):     {len(datos)} projects (2010-2018)",
        f"  Seguimiento (new):  {len(seg)} projects (2019-2025)",
        "",
        "Year distribution:",
    ]
    for yr, cnt in sorted(corpus["id"].str[:4].value_counts().sort_index().items()):
        lines.append(f"  {yr}: {cnt}")
    lines += [
        "",
        "Expert-labeled (gold set):",
        f"  {corpus['expert_category'].replace('', None).notna().sum()} projects",
        "",
        "Source breakdown:",
    ]
    for src, cnt in corpus["source"].value_counts().items():
        lines.append(f"  {src}: {cnt}")

    report = "\n".join(lines)
    OUT_STATS.write_text(report, encoding="utf-8")
    print(report)
    print(f"\nWrote corpus ({len(corpus)} rows) → {OUT_XLSX.relative_to(PAPER.parent.parent)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
